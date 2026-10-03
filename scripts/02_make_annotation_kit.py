"""
Step 2 — Build the blind label-audit kit (no user study; methodological annotation).

Samples 150 reviews stratified by rating-derived label (50/50/50), always including the
8 manual-override reviews. Star ratings, attraction names and existing labels are HIDDEN.
Two researchers annotate independently; scripts/03_label_agreement.py computes Cohen's kappa.
"""
import os

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
df = pd.read_csv(os.path.join(ROOT, "data/processed/reviews_publication_final.csv"))
SEED, N_PER = 2026, 50

over = df[df.Manual_Override == 1]
rest = df[df.Manual_Override == 0]
parts = [over]
for lab in ["Negative", "Neutral", "Positive"]:
    need = N_PER - (over.Rating_Label == lab).sum()
    parts.append(rest[rest.Rating_Label == lab].sample(need, random_state=SEED))
sample = pd.concat(parts).sample(frac=1, random_state=SEED).reset_index(drop=True)
sample["Item"] = range(1, len(sample) + 1)

os.makedirs(os.path.join(ROOT, "annotation"), exist_ok=True)
# key file (kept by the study lead, NOT given to annotators)
sample[["Item", "Review_ID", "Star_Rating", "Rating_Label", "Final_Sentiment", "Manual_Override"]].to_csv(
    os.path.join(ROOT, "annotation/annotation_key_DO_NOT_SHARE.csv"), index=False)

def sheet(annotator):
    wb = Workbook()
    ws = wb.active
    ws.title = "Annotate"
    ws.append(["Item", "Review text", "Label (Positive / Neutral / Negative)", "Confidence (1-3)", "Notes"])
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F4E79")
    for _, r in sample.iterrows():
        ws.append([int(r.Item), r.Original_Review_Text, "", "", ""])
    dv = DataValidation(type="list", formula1='"Positive,Neutral,Negative"', allow_blank=True)
    dv2 = DataValidation(type="list", formula1='"1,2,3"', allow_blank=True)
    ws.add_data_validation(dv); ws.add_data_validation(dv2)
    dv.add(f"C2:C{len(sample)+1}"); dv2.add(f"D2:D{len(sample)+1}")
    for col, w in zip("ABCDE", [6, 110, 22, 14, 30]):
        ws.column_dimensions[col].width = w
    for row in ws.iter_rows(min_row=2):
        row[1].alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A2"
    g = wb.create_sheet("Guidelines")
    for line in GUIDE.strip().splitlines():
        g.append([line])
    g.column_dimensions["A"].width = 120
    wb.save(os.path.join(ROOT, f"annotation/annotation_sheet_{annotator}.xlsx"))

GUIDE = """
LABEL-AUDIT GUIDELINES (read before annotating)
Judge ONLY the text. You do not see star ratings or attraction names on purpose.
Positive  - the reviewer's overall evaluation of the experience is favourable (praise dominates; would recommend).
Negative  - the overall evaluation is unfavourable (complaints dominate; would not recommend / warns others).
Neutral   - mixed with no clear dominant direction, mainly factual/informational, or explicitly 'average/ok'.
Mixed reviews: decide by the reviewer's overall verdict (often the final sentence or an explicit recommendation).
Sarcasm: label the intended meaning, not the literal words.
Confidence: 1 = unsure, 2 = fairly sure, 3 = certain.
Work independently. Do not discuss items with the other annotator until both sheets are submitted.
Do not look up the original review online.
"""

sheet("A")
sheet("B")
print(f"Sample: {len(sample)} items; rating-label mix:", sample.Rating_Label.value_counts().to_dict(),
      "; overrides included:", int(sample.Manual_Override.sum()))
